"""Load scraped Cal eProcure events + attachments into SQLite, extract attachment text,
and build a full-text index for search / retrieval.

    python ingest.py              # incremental: only new or changed attachments are re-read
    python ingest.py --rebuild    # wipe extracted text and redo everything
"""
import argparse
import glob
import io
import json
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from config import ATTACHMENTS_DIR, EVENTS_JSON, EXCEL_GLOB, PROJECT
from db import connect

MAX_PDF_PAGES = 150
CHUNK_CHARS = 1500


def parse_date(value):
    """'09/10/2026 4:09PM PDT' -> '2026-09-10T16:09' (Pacific local time, tz dropped)."""
    if not value:
        return None
    value = re.sub(r"\s+(PDT|PST|PT)$", "", value.strip())
    for fmt in ("%m/%d/%Y %I:%M%p", "%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M", "%m/%d/%Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%dT%H:%M")
        except ValueError:
            continue
    return None


def clean(text: str) -> str:
    text = text.replace("¿", "-").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def chunk(text: str, size: int = CHUNK_CHARS):
    paras, buf = re.split(r"\n\s*\n", text), ""
    for p in paras:
        if len(buf) + len(p) > size and buf:
            yield buf.strip()
            buf = ""
        while len(p) > size:
            yield p[:size]
            p = p[size:]
        buf += p + "\n\n"
    if buf.strip():
        yield buf.strip()


# ---------------------------------------------------------------- text extraction

def pdf_text(data: bytes):
    import pymupdf
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        pages = doc.page_count
        parts = [doc.load_page(i).get_text() for i in range(min(pages, MAX_PDF_PAGES))]
    return "\n".join(parts), pages


def docx_text(data: bytes):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts), None


def xlsx_text(data: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:8]:
        parts.append(f"## Sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 400:
                break
            vals = [str(v).strip() for v in row if v not in (None, "")]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts), None


def zip_text(data: bytes):
    parts = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist()[:60]:
            parts.append(f"## Zip member: {info.filename}")
            ext = info.filename.rsplit(".", 1)[-1].lower()
            if ext in EXTRACTORS and ext != "zip" and info.file_size < 40_000_000:
                try:
                    parts.append(EXTRACTORS[ext](zf.read(info))[0])
                except Exception as exc:  # one bad member must not sink the archive
                    parts.append(f"[unreadable: {exc}]")
    return "\n".join(parts), None


EXTRACTORS = {"pdf": pdf_text, "docx": docx_text, "xlsx": xlsx_text, "zip": zip_text}


def extract(path: Path):
    ext = path.suffix.lower().lstrip(".")
    if ext not in EXTRACTORS:
        return None, None, f"unsupported .{ext}"
    try:
        text, pages = EXTRACTORS[ext](path.read_bytes())
        return clean(text), pages, None
    except Exception as exc:
        return None, None, f"{type(exc).__name__}: {exc}"


# ---------------------------------------------------------------- excel exports

def read_exports():
    """Cal eProcure's 'Excel' export is really an HTML table. Newest file wins."""
    from bs4 import BeautifulSoup
    rows = {}
    for f in sorted(glob.glob(str(PROJECT / EXCEL_GLOB))):
        soup = BeautifulSoup(Path(f).read_text(encoding="utf-8", errors="ignore"), "html.parser")
        table = soup.find("table")
        if not table:
            continue
        trs = table.find_all("tr")
        headers = [th.get_text(strip=True) for th in trs[0].find_all(["th", "td"])]
        for tr in trs[1:]:
            cells = [td.get_text(strip=True) for td in tr.find_all("td")]
            if len(cells) == len(headers):
                rec = dict(zip(headers, cells))
                rows[rec.get("Event ID")] = rec
    return rows


# ---------------------------------------------------------------- main

def ingest(rebuild=False):
    conn = connect()
    now = datetime.now().isoformat(timespec="seconds")
    events = json.loads(EVENTS_JSON.read_text(encoding="utf-8"))
    exports = read_exports()
    print(f"{len(events)} events in JSON, {len(exports)} rows in Excel exports")

    if rebuild:
        conn.execute("DELETE FROM attachments")
    conn.execute("DELETE FROM chunks")
    conn.execute("DELETE FROM unspsc")

    seen = set()
    for ev in events:
        d = ev.get("Details") or {}
        eid = d.get("Event ID")
        if not eid:
            continue
        seen.add(eid)
        c = ev.get("Contact Information") or {}
        x = exports.get(eid, {})
        conn.execute(
            """INSERT INTO events (id,name,dept,dept_code,format,version,status,published,end_date,
                   description,contact_name,contact_email,contact_phone,detail_url,raw,first_seen,last_seen)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(id) DO UPDATE SET name=excluded.name, dept=excluded.dept,
                   dept_code=excluded.dept_code, format=excluded.format, version=excluded.version,
                   status=excluded.status, published=excluded.published, end_date=excluded.end_date,
                   description=excluded.description, contact_name=excluded.contact_name,
                   contact_email=excluded.contact_email, contact_phone=excluded.contact_phone,
                   detail_url=excluded.detail_url, raw=excluded.raw, last_seen=excluded.last_seen""",
            (eid, clean(ev.get("Event Name") or ""), d.get("Dept"), d.get("Department Code") or x.get("Department"),
             d.get("Format/Type"), d.get("Event Version"), d.get("Status") or x.get("Status"),
             parse_date(d.get("Published Date")), parse_date(d.get("Event End Date")) or parse_date(x.get("End Date")),
             clean(ev.get("Description") or "") or None, c.get("Contact Name") or x.get("Buyer Name"),
             c.get("Email") or x.get("Buyer Email"), c.get("Phone"), ev.get("Detail URL"),
             json.dumps(ev), now, now))

        for u in ev.get("UNSPSC Codes") or []:
            desc = u.get("UNSPSC Classification Description") or ""
            conn.execute("INSERT OR REPLACE INTO unspsc VALUES (?,?,?,?)",
                         (eid, u.get("UNSPSC Classification"), desc, desc.split(" - ")[0].strip()))

        if ev.get("Description"):
            for body in chunk(clean(ev["Description"])):
                conn.execute("INSERT INTO chunks VALUES (?,?,?)", (eid, "Description", body))

        for a in ev.get("Attachments") or []:
            fname = a.get("file_name")
            path = PROJECT / (a.get("local_path") or "").replace("\\", "/")
            ext = path.suffix.lower().lstrip(".")
            row = conn.execute("SELECT size_bytes, text FROM attachments WHERE event_id=? AND file_name=?",
                               (eid, fname)).fetchone()
            size = path.stat().st_size if path.is_file() else None
            if row and row["size_bytes"] == size and row["text"] is not None:
                text = row["text"]
            else:
                if path.is_file():
                    text, pages, err = extract(path)
                else:
                    text, pages, err = None, None, "file missing on disk"
                conn.execute("INSERT OR REPLACE INTO attachments VALUES (?,?,?,?,?,?,?,?,?,?)",
                             (eid, fname, str(path.relative_to(PROJECT)) if path.is_file() else a.get("local_path"),
                              ext, size, a.get("download_url"), pages, len(text or ""), text, err))
                n_chars = len(text or "")
                msg = ("ERR " + err) if err else f"{n_chars:>8} chars"
                print(f"  {eid}  {fname[:70]:70}  {msg}")
            for body in chunk(text or ""):
                conn.execute("INSERT INTO chunks VALUES (?,?,?)", (eid, fname, body))
        conn.commit()

    # Events that appear in the latest Excel export but were never scraped in detail.
    for eid, x in exports.items():
        if eid and eid not in seen:
            conn.execute(
                """INSERT OR IGNORE INTO events (id,name,dept,dept_code,format,status,end_date,
                       contact_name,contact_email,first_seen,last_seen) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (eid, x.get("Event Name"), x.get("Department Name"), x.get("Department"),
                 f"{x.get('Format', '')} / {x.get('Type', '')}", x.get("Status"), parse_date(x.get("End Date")),
                 x.get("Buyer Name"), x.get("Buyer Email"), now, now))
    conn.commit()

    s = conn.execute("""SELECT (SELECT COUNT(*) FROM events) e, (SELECT COUNT(*) FROM attachments) a,
                        (SELECT COUNT(*) FROM attachments WHERE error IS NOT NULL) ae,
                        (SELECT COUNT(*) FROM chunks) c""").fetchone()
    print(f"done: {s['e']} events, {s['a']} attachments ({s['ae']} unreadable), {s['c']} search chunks")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--rebuild", action="store_true")
    ingest(ap.parse_args().rebuild)
    sys.exit(0)
